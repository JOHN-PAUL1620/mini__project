import io
import tempfile
import unittest
from pathlib import Path

from docx import Document

from app import create_app
from database import certifications, internships


class ResumeAnalyzerTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.app = create_app({
            "TESTING": True,
            "DATABASE_URL": f"sqlite:///{(root / 'test.db').as_posix()}",
            "UPLOAD_FOLDER": str(root / "uploads"),
        })
        self.client = self.app.test_client()

    def tearDown(self):
        self.app.extensions["candidate_database"].engine.dispose()
        self.temp_dir.cleanup()

    def test_bulk_processing_persists_candidates_and_flags_duplicates(self):
        first_resume = _resume_bytes("Jane Developer", "jane@example.com")
        second_resume = _resume_bytes("John Analyst", "john@example.com")
        response = self.client.post("/api/resumes/analyze", data={
            "resumes": [(io.BytesIO(first_resume), "jane.docx"), (io.BytesIO(second_resume), "john.docx")]
        }, content_type="multipart/form-data")

        self.assertEqual(response.status_code, 200)
        results = response.get_json()["results"]
        self.assertEqual([item["status"] for item in results], ["processed", "processed"])
        self.assertIn("python", results[0]["candidate"]["skills"])
        self.assertNotIn("analysis", results[0]["candidate"])

        duplicate = self.client.post("/api/resumes/analyze", data={
            "resumes": (io.BytesIO(first_resume), "jane-copy.docx")
        }, content_type="multipart/form-data").get_json()["results"][0]
        self.assertEqual(duplicate["status"], "duplicate")
        self.assertEqual(duplicate["candidate_id"], results[0]["candidate_id"])

        candidates = self.client.get("/api/candidates").get_json()["candidates"]
        self.assertEqual(len(candidates), 2)
        self.assertFalse((Path(self.temp_dir.name) / "uploads").exists())

    def test_rejects_unsupported_file_type(self):
        response = self.client.post("/api/resumes/analyze", data={
            "resumes": (io.BytesIO(b"plain text"), "resume.txt")
        }, content_type="multipart/form-data")

        result = response.get_json()["results"][0]
        self.assertEqual(result["status"], "error")
        self.assertIn("PDF and DOCX", result["error"])

    def test_stores_partial_resume_with_missing_fields(self):
        response = self.client.post("/api/resumes/analyze", data={
            "resumes": (io.BytesIO(_resume_bytes("Minimal Candidate", "", include_details=False)), "minimal.docx")
        }, content_type="multipart/form-data")

        candidate = response.get_json()["results"][0]["candidate"]
        self.assertIn("email", candidate["missing_fields"])
        self.assertIn("education", candidate["missing_fields"])

    def test_delete_candidate_removes_complete_record(self):
        response = self.client.post("/api/resumes/analyze", data={
            "resumes": (io.BytesIO(_resume_bytes("Delete Candidate", "delete@example.com")), "delete.docx")
        }, content_type="multipart/form-data")
        candidate_id = response.get_json()["results"][0]["candidate_id"]

        delete_response = self.client.delete(f"/api/candidates/{candidate_id}")
        self.assertEqual(delete_response.status_code, 200)
        self.assertTrue(delete_response.get_json()["deleted"])
        self.assertEqual(self.client.get(f"/api/candidates/{candidate_id}").status_code, 404)
        self.assertEqual(self.client.delete(f"/api/candidates/{candidate_id}").status_code, 404)

        candidates = self.client.get("/api/candidates").get_json()["candidates"]
        self.assertEqual(candidates, [])

    def test_updates_candidate_profile(self):
        response = self.client.post("/api/resumes/analyze", data={
            "resumes": (io.BytesIO(_resume_bytes("Original Name", "original@example.com")), "original.docx")
        }, content_type="multipart/form-data")
        candidate_id = response.get_json()["results"][0]["candidate_id"]

        response = self.client.patch(f"/api/candidates/{candidate_id}", json={
            "full_name": "Updated Name", "email": "updated@example.com", "skills": ["python", "sql"],
            "education": ["B.Tech"], "work_experience": ["Developer"], "internship_experience": [],
        })

        self.assertEqual(response.status_code, 200)
        candidate = response.get_json()
        self.assertEqual(candidate["full_name"], "Updated Name")
        self.assertEqual(candidate["email"], "updated@example.com")
        self.assertEqual(candidate["skills"], ["python", "sql"])
        self.assertEqual(candidate["work_experience"], ["Developer"])

    def test_stores_certifications_and_internships_in_separate_records(self):
        response = self.client.post("/api/resumes/analyze", data={
            "resumes": (io.BytesIO(_resume_bytes("Records Candidate", "records@example.com", include_records=True)), "records.docx")
        }, content_type="multipart/form-data")
        candidate_id = response.get_json()["results"][0]["candidate_id"]

        with self.app.extensions["candidate_database"].engine.connect() as connection:
            saved_certifications = connection.execute(certifications.select().where(certifications.c.candidate_id == candidate_id)).mappings().all()
            saved_internships = connection.execute(internships.select().where(internships.c.candidate_id == candidate_id)).mappings().all()
        self.assertEqual(len(saved_certifications), 1)
        self.assertEqual(len(saved_internships), 1)

    def test_browser_flow_renders_processing_results(self):
        response = self.client.post("/analyze", data={
            "role": "Backend Developer",
            "resumes": (io.BytesIO(_resume_bytes("Browser Candidate", "browser@example.com")), "browser.docx")
        }, content_type="multipart/form-data")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Candidate results", response.data)
        self.assertIn(b"Matching skills", response.data)


def _resume_bytes(name, email, include_details=True, include_records=False):
    stream = io.BytesIO()
    document = Document()
    document.add_paragraph(name)
    if email:
        document.add_paragraph(f"{email} | +91 98765 43210")
    if include_details:
        for line in (
            "Skills", "Python, Flask, SQL, Docker, Git, Communication",
            "Experience", "Backend Engineer - Example Ltd - 2022 to 2025",
            "Education", "B.Tech Computer Science - Example University - 2022 - CGPA: 8.8",
            "Projects", "Resume Intelligence API built with Python and Flask",
        ):
            document.add_paragraph(line)
    if include_records:
        for line in (
            "Certifications", "AWS Certified Cloud Practitioner",
            "Internships", "Software Intern - Example Ltd - 2021",
        ):
            document.add_paragraph(line)
    document.save(stream)
    return stream.getvalue()


if __name__ == "__main__":
    unittest.main()
