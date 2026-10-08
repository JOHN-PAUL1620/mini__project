import io
import unittest

from docx import Document

from parser import parse_resume


class ParserTest(unittest.TestCase):
    def test_collects_a_single_word_name_from_the_first_line(self):
        stream = io.BytesIO()
        document = Document()
        document.add_paragraph("HITHESH")
        document.add_paragraph("Data Analyst")
        document.add_paragraph("hithesh@example.com")
        document.add_paragraph("PROFESSIONAL SUMMARY")
        document.add_paragraph("Analyzes business data.")
        document.save(stream)

        parsed = parse_resume(stream.getvalue(), "hithesh.docx")

        self.assertEqual(parsed["full_name"], "Hithesh")

    def test_collects_linkedin_without_protocol(self):
        stream = io.BytesIO()
        document = Document()
        document.add_paragraph("Jane Developer")
        document.add_paragraph("jane@example.com | +91 98765 43210")
        document.add_paragraph("LinkedIn: linkedin.com/in/jane-developer")
        document.add_paragraph("Skills")
        document.add_paragraph("Python, Flask, SQL")
        document.save(stream)

        parsed = parse_resume(stream.getvalue(), "jane.docx")

        self.assertEqual(parsed["linkedin"], "https://linkedin.com/in/jane-developer")

    def test_collects_labeled_linkedin_and_github_placeholder_urls(self):
        stream = io.BytesIO()
        document = Document()
        document.add_paragraph("Sophia Mitchell")
        document.add_paragraph("Email: sample@example-resume.test")
        document.add_paragraph("Phone: +1 555-123-4567")
        document.add_paragraph("LinkedIn: linkedin.example.com/in/sample-profile")
        document.add_paragraph("GitHub: github.example.com/sample-profile")
        document.add_paragraph("Skills")
        document.add_paragraph("Python, Communication, Teamwork, Problem Solving")
        document.save(stream)

        parsed = parse_resume(stream.getvalue(), "sophia.docx")

        self.assertEqual(parsed["linkedin"], "https://linkedin.example.com/in/sample-profile")
        self.assertEqual(parsed["github"], "https://github.example.com/sample-profile")


if __name__ == "__main__":
    unittest.main()
