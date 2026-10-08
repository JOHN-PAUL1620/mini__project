"""Resume text extraction and lightweight structured parsing."""

from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

SKILLS = {
    "python", "java", "javascript", "typescript", "react", "angular", "vue",
    "node.js", "flask", "django", "fastapi", "spring", "sql", "mysql",
    "postgresql", "mongodb", "redis", "html", "css", "aws", "azure", "gcp",
    "docker", "kubernetes", "git", "linux", "machine learning", "deep learning",
    "tensorflow", "pytorch", "pandas", "numpy", "power bi", "tableau", "excel",
    "communication", "leadership", "teamwork", "problem solving",
}

SOFT_SKILLS = {"communication", "leadership", "teamwork", "problem solving"}

SECTION_ALIASES = {
    "education": ("education", "academic background", "qualification"),
    "certifications": ("certifications", "certificates", "licenses"),
    "experience": ("experience", "work experience", "employment history"),
    "internships": ("internships", "internship experience"),
    "projects": ("projects", "personal projects", "academic projects"),
    "achievements": ("achievements", "awards", "accomplishments"),
    "languages": ("languages", "languages known"),
}


class ResumeParseError(ValueError):
    """Raised when a resume cannot be read."""


def extract_text(file_bytes: bytes, filename: str) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError("Only PDF and DOCX resumes are supported.")

    try:
        if extension == ".docx":
            document = Document(io.BytesIO(file_bytes))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            table_text = [
                " ".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
            text = "\n".join([text, *table_text])
        else:
            import pdfplumber

            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as exc:
        raise ResumeParseError(f"Could not read {extension[1:].upper()} resume.") from exc

    text = re.sub(r"\r\n?", "\n", text).strip()
    if not text:
        raise ResumeParseError("The resume does not contain extractable text.")
    return text


def parse_resume(file_or_bytes, filename: str | None = None) -> dict:
    """Return structured fields from a resume upload or raw bytes."""
    if hasattr(file_or_bytes, "read"):
        filename = filename or getattr(file_or_bytes, "filename", "")
        file_bytes = file_or_bytes.read()
    else:
        file_bytes = file_or_bytes
        filename = filename or ""

    text = extract_text(file_bytes, filename)
    lines = [line.strip(" \t|") for line in text.splitlines() if line.strip()]
    lower_text = text.lower()
    email = _first(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)
    phone = _first(r"(?:\+?\d[\d\s().-]{8,}\d)", text)
    phone = re.sub(r"[^\d+]", "", phone) if phone else ""
    skills = sorted(skill for skill in SKILLS if _contains_skill(lower_text, skill))

    sections = {name: _section(lines, aliases) for name, aliases in SECTION_ALIASES.items()}
    education_text = sections["education"]
    missing_fields = []
    result = {
        "full_name": _candidate_name(lines),
        "email": email,
        "phone": phone,
        "linkedin": _profile_url("linkedin", r"(?:https?://)?(?:www\.)?linkedin\.com/[^\s,]+", text),
        "github": _profile_url("github", r"(?:https?://)?(?:www\.)?github\.com/[^\s,]+", text),
        "portfolio": _portfolio(text),
        "address": _address(lines),
        "skills": skills,
        "technical_skills": [skill for skill in skills if skill not in SOFT_SKILLS],
        "soft_skills": [skill for skill in skills if skill in SOFT_SKILLS],
        "certifications": _as_items(sections["certifications"]),
        "work_experience": _as_items(sections["experience"]),
        "internship_experience": _as_items(sections["internships"]),
        "projects": _as_items(sections["projects"]),
        "education": _as_items(education_text),
        "cgpa_percentage": _first(r"\b(?:CGPA|GPA|percentage|score)\s*[:\-]?\s*([0-9.]+\s*%?)", text, 1),
        "college_name": _college_name(education_text),
        "achievements": _as_items(sections["achievements"]),
        "languages_known": _languages(sections["languages"]),
        "total_years_experience": _years_of_experience(sections["experience"]),
        "raw_text": text,
    }
    for field in ("full_name", "email", "phone", "skills", "education"):
        if not result[field]:
            missing_fields.append(field)
    result["missing_fields"] = missing_fields
    return result


def _first(pattern: str, text: str, group: int = 0) -> str:
    match = re.search(pattern, text, re.IGNORECASE)
    return match.group(group).strip(" .,;") if match else ""


def _profile_url(label: str, pattern: str, text: str) -> str:
    url = _first(pattern, text) or _labeled_url(label, text)
    if url and not re.match(r"https?://", url, re.I):
        return f"https://{url}"
    return url


def _labeled_url(label: str, text: str) -> str:
    pattern = rf"\b{re.escape(label)}\s*[:\-]\s*(?:https?://)?([^\s,|;]+)"
    match = re.search(pattern, text, re.IGNORECASE)
    return match.group(1).strip(" .,)") if match else ""


def _contains_skill(text: str, skill: str) -> bool:
    return bool(re.search(rf"(?<!\w){re.escape(skill)}(?!\w)", text))


def _candidate_name(lines: list[str]) -> str:
    for index, line in enumerate(lines[:6]):
        word_count = len(line.split())
        if (
            word_count in range(2, 5)
            and len(line) < 60
            and not re.search(r"@|https?://|\d|resume|curriculum|linkedin|github", line, re.I)
        ):
            return line.title() if line.isupper() else line
        # Some resumes put a single given name on the first line. Restrict this
        # case to the first line so later headings cannot become the candidate name.
        if (
            index == 0
            and word_count == 1
            and len(line) < 60
            and re.fullmatch(r"[A-Za-z][A-Za-z.'-]*", line)
            and not re.search(r"resume|curriculum", line, re.I)
        ):
            return line.title() if line.isupper() else line
    return ""


def _section(lines: list[str], aliases: tuple[str, ...]) -> str:
    all_headings = {alias for values in SECTION_ALIASES.values() for alias in values}
    start = None
    collected = []
    for line in lines:
        normalized = re.sub(r"[^a-z ]", "", line.lower()).strip()
        is_heading = normalized in all_headings or normalized.rstrip("s") in all_headings
        if start is not None and is_heading:
            break
        if normalized in aliases:
            start = True
            continue
        if start:
            collected.append(line)
    return "\n".join(collected)


def _as_items(text: str) -> list[str]:
    return [line.lstrip("-*• ").strip() for line in text.splitlines() if line.strip()][:20]


def _address(lines: list[str]) -> str:
    return next((line for line in lines if re.search(r"\baddress\s*:", line, re.I)), "")


def _portfolio(text: str) -> str:
    urls = re.findall(r"https?://[^\s,]+", text, re.I)
    return next((url.strip(".,;)") for url in urls if "linkedin.com" not in url and "github.com" not in url), "")


def _college_name(education: str) -> str:
    return next((line for line in education.splitlines() if re.search(r"college|university|institute", line, re.I)), "")


def _languages(text: str) -> list[str]:
    return [item.strip() for item in re.split(r"[,|/]", text.replace("\n", ",")) if item.strip()][:10]


def _years_of_experience(experience_text: str) -> float:
    explicit = re.search(r"(\d+(?:\.\d+)?)\+?\s+years?(?:\s+of)?\s+experience", experience_text, re.I)
    if explicit:
        return float(explicit.group(1))

    years = [int(year) for year in re.findall(r"\b(?:19|20)\d{2}\b", experience_text)]
    current_year = datetime.now().year
    valid_years = [year for year in years if 1970 <= year <= current_year]
    return float(min(current_year - min(valid_years), 50)) if valid_years else 0.0
